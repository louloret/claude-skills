#!/usr/bin/env python3
"""Convert a Luis-style newsletter markdown to a LinkedIn-ready Word document."""

import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Helpers ──────────────────────────────────────────────────────────────────

def add_left_border(paragraph, color="AAAAAA", size=18, space=10):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(size))
    left.set(qn('w:space'), str(space))
    left.set(qn('w:color'), color)
    pBdr.append(left)
    pPr.append(pBdr)


def set_para_shading(paragraph, fill="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def set_cell_bg(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def add_inline_runs(paragraph, text, size=Pt(11), font='Georgia',
                    color=RGBColor(0x1A, 0x1A, 0x1A), bold=False, italic=False):
    """Add runs with **bold** and *italic* inline handling."""
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*)'
    for part in re.split(pattern, text):
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
        run.font.name = font
        run.font.size = size
        run.font.color.rgb = color


# ── Main converter ────────────────────────────────────────────────────────────

def convert(md_path: Path, docx_path: Path):
    doc = Document()

    # Page layout
    sec = doc.sections[0]
    sec.page_width  = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(1.25)
    sec.top_margin  = sec.bottom_margin = Inches(1.0)

    # Base style
    normal = doc.styles['Normal']
    normal.font.name = 'Georgia'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    lines = md_path.read_text(encoding='utf-8').split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]

        # ── H1: Post title ───────────────────────────────────────────────────
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.name  = 'Georgia'
            run.font.size  = Pt(26)
            run.font.bold  = True
            run.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0D)
            i += 1
            continue

        # ── H2: Author name ──────────────────────────────────────────────────
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(1)
            run = p.add_run(line[3:].strip())
            run.font.name  = 'Georgia'
            run.font.size  = Pt(11)
            run.font.bold  = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue

        # ── H3: Section header ───────────────────────────────────────────────
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(22)
            p.paragraph_format.space_after  = Pt(6)
            run = p.add_run(line[4:].strip())
            run.font.name  = 'Georgia'
            run.font.size  = Pt(14)
            run.font.bold  = True
            run.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0D)
            i += 1
            continue

        # ── Code block (monospace diagrams) ──────────────────────────────────
        if line.startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent  = Inches(0.2)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after  = Pt(1)
                set_para_shading(p, 'F2F2F2')
                run = p.add_run(cl if cl else ' ')
                run.font.name  = 'Courier New'
                run.font.size  = Pt(9)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            continue

        # ── Markdown table ───────────────────────────────────────────────────
        if line.startswith('|'):
            raw_rows = []
            while i < len(lines) and lines[i].startswith('|'):
                raw_rows.append(lines[i])
                i += 1
            rows = []
            for r in raw_rows:
                if re.match(r'^\|[\s\-:|]+\|$', r):
                    continue  # separator row
                rows.append([c.strip() for c in r.strip('|').split('|')])
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncols)
            tbl.style = 'Table Grid'
            for r_idx, cells in enumerate(rows):
                for c_idx in range(ncols):
                    cell = tbl.rows[r_idx].cells[c_idx]
                    text = cells[c_idx] if c_idx < len(cells) else ''
                    cell.text = ''
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after  = Pt(3)
                    run = p.add_run(text)
                    run.font.name = 'Georgia'
                    run.font.size = Pt(10)
                    if r_idx == 0:
                        run.font.bold = True
                        set_cell_bg(cell, 'EBEBEB')
            doc.add_paragraph()
            continue

        # ── Blockquote ───────────────────────────────────────────────────────
        if line.startswith('>'):
            bq_lines = []
            while i < len(lines) and lines[i].startswith('>'):
                bq_lines.append(lines[i])
                i += 1
            for bq in bq_lines:
                text = bq.lstrip('>').lstrip(' ')
                p = doc.add_paragraph()
                p.paragraph_format.left_indent  = Inches(0.35)
                p.paragraph_format.right_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after  = Pt(3)
                add_left_border(p)
                if text.strip():
                    add_inline_runs(p, text.rstrip(),
                                    size=Pt(10.5),
                                    color=RGBColor(0x33, 0x33, 0x33))
                else:
                    p.add_run(' ').font.size = Pt(4)
            continue

        # ── Arrow connector lines (↓ between numbered steps) ─────────────────
        if re.match(r'^\s+[↓↑→←]\s*$', line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent  = Inches(0.25)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(0)
            run = p.add_run(line.strip())
            run.font.name  = 'Georgia'
            run.font.size  = Pt(10)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            i += 1
            continue

        # ── Empty line ───────────────────────────────────────────────────────
        if not line.strip():
            i += 1
            continue

        # ── Regular paragraph ─────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(6)
        add_inline_runs(p, line)
        i += 1

    doc.save(str(docx_path))
    print(f"Saved: {docx_path}")


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 md_to_linkedin_docx.py input.md [output.docx]")
        sys.exit(1)
    md   = Path(sys.argv[1])
    out  = Path(sys.argv[2]) if len(sys.argv) >= 3 else md.with_suffix('.docx')
    convert(md, out)
